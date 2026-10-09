"""Single-file RAG pipeline over ./data using Azure OpenAI.

파이프라인: 수집(Load) → 전처리(Preprocess) → 분할(Split) → 임베딩/저장(Embed & Store)
          → 검색(Retrieve) → 생성(Generate)

지원 형식: .md / .txt / .pdf / .csv / .docx
분할 전략: fixed / recursive / sentence / markdown  (--splitter 로 선택, --experiment 로 비교)

Usage:
    pip install -r requirements.txt
    python main.py                                  # 대화형 Q&A (기본 recursive 분할)
    python main.py "Gatik의 핵심 기술은?"              # 단발 질의
    python main.py --splitter sentence "질문"        # 분할 전략 지정
    python main.py --retriever mmr "질문"            # MMR 검색 (다양성 고려)
    python main.py --experiment                     # 분할 전략별 검색 성능 비교
    python main.py --rebuild                        # 임베딩 캐시 무시하고 재생성

Required env vars:
    AOAI_ENDPOINT, AOAI_API_KEY,
    AOAI_DEPLOY_GPT4O_MINI (chat), AOAI_DEPLOY_EMBED_3_SMALL (embedding)
Optional:
    AOAI_API_VERSION (default 2024-10-21)
    RAG_CHAT_DEPLOY / RAG_EMBED_DEPLOY 로 배포 이름 override
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import sys
import unicodedata
from dataclasses import asdict, dataclass, field
from pathlib import Path

import numpy as np
from openai import AzureOpenAI

ROOT = Path(__file__).resolve().parent


# =====================================================================
# 0. Config
# =====================================================================
@dataclass
class Config:
    data_dir: Path = ROOT / "data"
    cache_dir: Path = ROOT / ".rag_cache"
    splitter: str = "recursive"
    chunk_size: int = 800
    chunk_overlap: int = 120
    top_k: int = 5
    retriever: str = "similarity"  # similarity | mmr
    mmr_lambda: float = 0.5
    embed_batch: int = 64
    temperature: float = 0.2


def require_env(name: str) -> str:
    value = os.getenv(name)
    if not value:
        sys.exit(f"[error] 환경변수 {name} 가 설정되지 않았습니다.")
    return value


def make_client() -> AzureOpenAI:
    return AzureOpenAI(
        azure_endpoint=require_env("AOAI_ENDPOINT"),
        api_key=require_env("AOAI_API_KEY"),
        api_version=os.getenv("AOAI_API_VERSION", "2024-10-21"),
    )


@dataclass
class Document:
    text: str
    metadata: dict = field(default_factory=dict)


# =====================================================================
# 1. 수집 (Load): 형식별 로더
# =====================================================================
def load_text(path: Path) -> list[Document]:
    return [Document(path.read_text(encoding="utf-8", errors="ignore"), {"type": path.suffix[1:]})]


def load_pdf(path: Path) -> list[Document]:
    """페이지 단위로 Document 생성 → 출처에 페이지 번호를 남길 수 있음."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return [
        Document(page.extract_text() or "", {"type": "pdf", "page": i})
        for i, page in enumerate(reader.pages, 1)
    ]


def load_csv(path: Path) -> list[Document]:
    """행 하나를 '컬럼: 값' 형태의 자연어 레코드로 변환 (표 데이터의 의미 보존)."""
    with path.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    return [
        Document(
            "\n".join(f"{k}: {v}" for k, v in row.items() if v),
            {"type": "csv", "row": i},
        )
        for i, row in enumerate(rows, 1)
    ]


def load_docx(path: Path) -> list[Document]:
    """본문 문단 + 표(행을 ' | '로 연결)를 함께 추출."""
    import docx

    d = docx.Document(str(path))
    parts = []
    for p in d.paragraphs:
        if not p.text.strip():
            continue
        # Heading 스타일은 markdown 헤더로 변환 → markdown 분할기에서 활용
        if p.style.name.startswith("Heading"):
            level = "".join(c for c in p.style.name if c.isdigit()) or "1"
            parts.append("#" * int(level) + " " + p.text)
        else:
            parts.append(p.text)
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [Document("\n".join(parts), {"type": "docx"})]


LOADERS = {
    ".md": load_text,
    ".txt": load_text,
    ".pdf": load_pdf,
    ".csv": load_csv,
    ".docx": load_docx,
}


def load_documents(data_dir: Path) -> list[Document]:
    docs = []
    for path in sorted(data_dir.rglob("*")):
        loader = LOADERS.get(path.suffix.lower())
        if not path.is_file() or loader is None:
            continue
        try:
            loaded = loader(path)
        except Exception as e:  # 깨진/암호화된 파일 등
            print(f"[warn] 로드 실패 {path}: {e}")
            continue
        rel = str(path.relative_to(data_dir))
        for doc in loaded:
            doc.metadata.update(source=rel, category=path.parent.name)
        docs.extend(loaded)
    return docs


# =====================================================================
# 2. 전처리 (Preprocess)
# =====================================================================
def clean_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\x00", "")
    text = re.sub(r"-\n(?=\w)", "", text)  # PDF 줄바꿈 하이픈 결합
    text = re.sub(r"[ \t]+", " ", text)  # 연속 공백 축소
    text = re.sub(r"\n\s*\n+", "\n\n", text)  # 빈 줄 정리 (문단 경계는 유지)
    text = re.sub(r"^\s*\d+\s*$", "", text, flags=re.M)  # 페이지 번호만 있는 줄 제거
    return text.strip()


def preprocess(docs: list[Document], min_chars: int = 30) -> list[Document]:
    seen, out = set(), []
    for doc in docs:
        text = clean_text(doc.text)
        digest = hashlib.md5(text.encode()).hexdigest()
        if len(text) < min_chars or digest in seen:  # 너무 짧거나 중복인 문서 제거
            continue
        seen.add(digest)
        out.append(Document(text, doc.metadata))
    return out


# =====================================================================
# 3. 분할 (Split): 전략별 분할기
# =====================================================================
def split_fixed(text: str, size: int, overlap: int) -> list[str]:
    """고정 길이 문자 윈도우. 가장 단순하지만 문장/단어 중간에서 끊길 수 있음."""
    step = max(size - overlap, 1)
    return [text[i : i + size] for i in range(0, len(text), step)]


def _merge(pieces: list[str], size: int, overlap: int, sep: str) -> list[str]:
    """작은 조각들을 size 이하로 합치고, 이전 청크 끝부분을 overlap 만큼 이어붙임."""
    chunks, cur = [], ""
    for piece in pieces:
        if len(piece) > size:  # 단일 조각이 너무 길면 고정 길이로 강제 분할
            if cur:
                chunks.append(cur)
                cur = ""
            chunks.extend(split_fixed(piece, size, overlap))
            continue
        candidate = f"{cur}{sep}{piece}" if cur else piece
        if len(candidate) <= size:
            cur = candidate
            continue
        chunks.append(cur)
        tail = cur[-overlap:] if overlap else ""
        with_tail = f"{tail}{sep}{piece}"
        cur = with_tail if tail and len(with_tail) <= size else piece
    if cur.strip():
        chunks.append(cur)
    return chunks


def split_recursive(text: str, size: int, overlap: int,
                    seps: tuple[str, ...] = ("\n\n", "\n", ". ", " ")) -> list[str]:
    """문단 → 줄 → 문장 → 단어 순으로 큰 경계부터 시도 (LangChain RecursiveCharacterTextSplitter 방식)."""
    if len(text) <= size:
        return [text]
    for i, sep in enumerate(seps):
        if sep in text:
            pieces = []
            for p in text.split(sep):
                if len(p) > size:
                    pieces.extend(split_recursive(p, size, 0, seps[i + 1 :]))
                elif p.strip():
                    pieces.append(p)
            return _merge(pieces, size, overlap, sep)
    return split_fixed(text, size, overlap)


SENTENCE_END = re.compile(r"(?<=[.!?。])\s+|\n+")


def split_sentence(text: str, size: int, overlap: int) -> list[str]:
    """문장 단위로 자른 뒤 size 까지 묶음. 의미 단위 보존에 유리."""
    sentences = [s.strip() for s in SENTENCE_END.split(text) if s and s.strip()]
    return _merge(sentences, size, overlap, " ")


def split_markdown(text: str, size: int, overlap: int) -> list[str]:
    """# 헤더 기준 섹션 분할 후, 섹션 제목을 각 청크 앞에 붙여 문맥 유지."""
    sections = re.split(r"(?m)^(?=#{1,6} )", text)
    chunks = []
    for sec in sections:
        if not sec.strip():
            continue
        header = sec.splitlines()[0] if sec.startswith("#") else ""
        for c in split_recursive(sec, size, overlap):
            chunks.append(c if not header or c.startswith(header) else f"{header}\n{c}")
    return chunks


SPLITTERS = {
    "fixed": split_fixed,
    "recursive": split_recursive,
    "sentence": split_sentence,
    "markdown": split_markdown,
}


def split_documents(docs: list[Document], cfg: Config) -> list[Document]:
    splitter = SPLITTERS[cfg.splitter]
    chunks = []
    for doc in docs:
        for i, text in enumerate(splitter(doc.text, cfg.chunk_size, cfg.chunk_overlap)):
            if text.strip():
                chunks.append(Document(text.strip(), {**doc.metadata, "chunk": i}))
    return chunks


# =====================================================================
# 4. 임베딩 & 벡터 저장소 (Embed & Store)
# =====================================================================
class Embedder:
    def __init__(self, client: AzureOpenAI, deployment: str, batch: int):
        self.client, self.deployment, self.batch = client, deployment, batch

    def __call__(self, texts: list[str]) -> np.ndarray:
        vectors = []
        for i in range(0, len(texts), self.batch):
            resp = self.client.embeddings.create(model=self.deployment, input=texts[i : i + self.batch])
            vectors.extend(d.embedding for d in resp.data)
        arr = np.asarray(vectors, dtype=np.float32)
        return arr / np.linalg.norm(arr, axis=1, keepdims=True)  # 정규화 → 내적 = 코사인 유사도


class VectorStore:
    """numpy 기반 인메모리 벡터 저장소 (디스크 캐시 지원)."""

    def __init__(self, docs: list[Document], vectors: np.ndarray):
        self.docs, self.vectors = docs, vectors

    @classmethod
    def build(cls, docs: list[Document], embedder: Embedder, cache_dir: Path,
              rebuild: bool = False) -> "VectorStore":
        key = hashlib.sha256(
            (embedder.deployment + json.dumps([asdict(d) for d in docs], ensure_ascii=False)).encode()
        ).hexdigest()[:16]
        meta_path, vec_path = cache_dir / f"{key}.json", cache_dir / f"{key}.npy"
        if not rebuild and meta_path.exists() and vec_path.exists():
            print(f"[store] 캐시 사용 ({len(docs)} chunks)")
            return cls(docs, np.load(vec_path))

        print(f"[store] {len(docs)} chunks 임베딩 중...")
        vectors = embedder([d.text for d in docs])
        cache_dir.mkdir(exist_ok=True)
        meta_path.write_text(json.dumps([asdict(d) for d in docs], ensure_ascii=False), encoding="utf-8")
        np.save(vec_path, vectors)
        return cls(docs, vectors)

    def scores(self, query_vec: np.ndarray) -> np.ndarray:
        return self.vectors @ query_vec


# =====================================================================
# 5. 검색 (Retrieve)
# =====================================================================
class Retriever:
    def __init__(self, store: VectorStore, embedder: Embedder, cfg: Config):
        self.store, self.embedder, self.cfg = store, embedder, cfg

    def __call__(self, query: str) -> list[tuple[Document, float]]:
        q = self.embedder([query])[0]
        scores = self.store.scores(q)
        if self.cfg.retriever == "mmr":
            idx = self._mmr(scores)
        else:
            idx = np.argsort(-scores)[: self.cfg.top_k]
        return [(self.store.docs[i], float(scores[i])) for i in idx]

    def _mmr(self, scores: np.ndarray, fetch_k: int = 20) -> list[int]:
        """Maximal Marginal Relevance: 관련성은 높고 서로 중복은 적은 청크 선택."""
        candidates = list(np.argsort(-scores)[:fetch_k])
        vecs = self.store.vectors
        selected: list[int] = []
        while candidates and len(selected) < self.cfg.top_k:
            if selected:
                redundancy = (vecs[candidates] @ vecs[selected].T).max(axis=1)
            else:
                redundancy = np.zeros(len(candidates))
            mmr = self.cfg.mmr_lambda * scores[candidates] - (1 - self.cfg.mmr_lambda) * redundancy
            best = candidates.pop(int(np.argmax(mmr)))
            selected.append(best)
        return selected


# =====================================================================
# 6. 생성 (Generate)
# =====================================================================
SYSTEM_PROMPT = (
    "당신은 물류/유통 AI 스타트업 투자 분석가입니다. "
    "반드시 제공된 [컨텍스트]에 근거해서 한국어로 답하세요. "
    "근거가 없으면 '자료에서 확인되지 않습니다'라고 답하고, "
    "각 주장 뒤에 사용한 출처 번호를 [1], [2] 형태로 표시하세요."
)


def cite(doc: Document) -> str:
    m = doc.metadata
    loc = f" p.{m['page']}" if "page" in m else f" row {m['row']}" if "row" in m else ""
    return f"{m['source']}{loc}"


class Generator:
    def __init__(self, client: AzureOpenAI, deployment: str, cfg: Config):
        self.client, self.deployment, self.cfg = client, deployment, cfg

    def __call__(self, query: str, hits: list[tuple[Document, float]]) -> str:
        context = "\n\n".join(f"[{i}] ({cite(d)})\n{d.text}" for i, (d, _) in enumerate(hits, 1))
        resp = self.client.chat.completions.create(
            model=self.deployment,
            temperature=self.cfg.temperature,
            messages=[
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user", "content": f"[컨텍스트]\n{context}\n\n[질문]\n{query}"},
            ],
        )
        return resp.choices[0].message.content


# =====================================================================
# 7. 파이프라인 조립
# =====================================================================
class RAGPipeline:
    def __init__(self, cfg: Config, rebuild: bool = False, docs: list[Document] | None = None):
        client = make_client()
        self.cfg = cfg
        self.embedder = Embedder(
            client, os.getenv("RAG_EMBED_DEPLOY") or require_env("AOAI_DEPLOY_EMBED_3_SMALL"), cfg.embed_batch
        )
        self.generator = Generator(
            client, os.getenv("RAG_CHAT_DEPLOY") or require_env("AOAI_DEPLOY_GPT4O_MINI"), cfg
        )

        if docs is None:
            docs = prepare_documents(cfg)
        self.chunks = split_documents(docs, cfg)
        print(f"[split] strategy={cfg.splitter} → {len(self.chunks)} chunks")
        self.store = VectorStore.build(self.chunks, self.embedder, cfg.cache_dir, rebuild)
        self.retriever = Retriever(self.store, self.embedder, cfg)

    def ask(self, query: str) -> str:
        hits = self.retriever(query)
        answer = self.generator(query, hits)
        sources = "\n".join(f"  [{i}] {cite(d)} (score={s:.3f})" for i, (d, s) in enumerate(hits, 1))
        return f"{answer}\n\n출처:\n{sources}"


def prepare_documents(cfg: Config) -> list[Document]:
    raw = load_documents(cfg.data_dir)
    docs = preprocess(raw)
    by_type: dict[str, int] = {}
    for d in docs:
        by_type[d.metadata["type"]] = by_type.get(d.metadata["type"], 0) + 1
    print(f"[load] {len(raw)} raw → {len(docs)} cleaned docs {by_type}")
    if not docs:
        sys.exit(f"[error] {cfg.data_dir} 에 읽을 문서가 없습니다.")
    return docs


# =====================================================================
# 8. 분할 전략 실험 (Experiment)
# =====================================================================
# (질문, 정답 문서 경로에 포함되어야 할 키워드)
EVAL_SET = [
    ("Gatik의 자율주행 미들마일 서비스는 어떤 고객을 대상으로 하나?", "gatik"),
    ("GreyOrange 창고 오케스트레이션의 핵심 기술은?", "greyorange"),
    ("Flexport 플랫폼은 어떤 기능을 제공하나?", "flexport"),
    ("물류 AI 시장 규모와 성장률은?", "market"),
    ("Gatik, Flexport, GreyOrange의 강점과 약점 비교", "competitors"),
    ("AI 물류 스타트업 후보 목록", "scout"),
]


def run_experiment(base: Config, rebuild: bool) -> None:
    docs = prepare_documents(base)
    rows = []
    for name in SPLITTERS:
        cfg = Config(**{**asdict(base), "splitter": name})
        pipe = RAGPipeline(cfg, rebuild, docs=docs)
        lengths = [len(c.text) for c in pipe.chunks]
        hit, rr, top1 = 0, 0.0, []
        for query, expected in EVAL_SET:
            results = pipe.retriever(query)
            top1.append(results[0][1])
            ranks = [i for i, (d, _) in enumerate(results, 1) if expected in d.metadata["source"].lower()]
            if ranks:
                hit += 1
                rr += 1 / ranks[0]
        n = len(EVAL_SET)
        rows.append((name, len(lengths), np.mean(lengths), hit / n, rr / n, np.mean(top1)))

    print(f"\n=== 분할 전략 비교 (chunk_size={base.chunk_size}, overlap={base.chunk_overlap}, top_k={base.top_k}) ===")
    print(f"{'splitter':<10} {'chunks':>7} {'avg_len':>8} {'hit@k':>7} {'MRR':>6} {'top1_sim':>9}")
    for name, cnt, avg, hit_rate, mrr, sim in rows:
        print(f"{name:<10} {cnt:>7} {avg:>8.0f} {hit_rate:>7.2f} {mrr:>6.2f} {sim:>9.3f}")
    best = max(rows, key=lambda r: (r[4], r[5]))
    print(f"\n→ MRR 기준 최적 전략: {best[0]}")


# =====================================================================
# CLI
# =====================================================================
def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Azure OpenAI 기반 단일 파일 RAG")
    p.add_argument("query", nargs="*", help="질문 (없으면 대화형 모드)")
    p.add_argument("--splitter", choices=SPLITTERS, default="recursive")
    p.add_argument("--chunk-size", type=int, default=800)
    p.add_argument("--chunk-overlap", type=int, default=120)
    p.add_argument("--top-k", type=int, default=5)
    p.add_argument("--retriever", choices=["similarity", "mmr"], default="similarity")
    p.add_argument("--data-dir", type=Path, default=ROOT / "data")
    p.add_argument("--experiment", action="store_true", help="분할 전략별 검색 성능 비교")
    p.add_argument("--rebuild", action="store_true", help="임베딩 캐시 무시")
    return p.parse_args()


def main() -> None:
    args = parse_args()
    cfg = Config(
        data_dir=args.data_dir,
        splitter=args.splitter,
        chunk_size=args.chunk_size,
        chunk_overlap=args.chunk_overlap,
        top_k=args.top_k,
        retriever=args.retriever,
    )

    if args.experiment:
        run_experiment(cfg, args.rebuild)
        return

    pipe = RAGPipeline(cfg, rebuild=args.rebuild)
    if args.query:
        print(pipe.ask(" ".join(args.query)))
        return

    print("질문을 입력하세요 (종료: exit / Ctrl+D)")
    while True:
        try:
            query = input("\n> ").strip()
        except (EOFError, KeyboardInterrupt):
            print()
            break
        if query.lower() in {"exit", "quit", "q"}:
            break
        if query:
            print(pipe.ask(query))


if __name__ == "__main__":
    main()
